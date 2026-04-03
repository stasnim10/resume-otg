from docx import Document
import json

# Load both documents
original = Document('/Users/simum/Personal/My Mac/Documents/Simon Business School/Job Hunt/2026 Full Time Roles/USA/Draft/Simum Tasnim_Resume_Draft.docx')
optimized = Document('/Users/simum/Personal/My Mac/Documents/Simon Business School/Job Hunt/2026 Full Time Roles/USA/Draft/Simum Tasnim_Resume_Draft_Optimized.docx')

# The JSON payload from user
json_payload = """{
"summary_replacement": {
"match_anchor": "MBA candidate and strategy-driven operations professional with experience supporting revenue, growth, and go-to-market execution",
"replacement_text": "MBA candidate (Class of 2026) with 6+ years of international supply chain and logistics leadership, managing $50M–$240M global operations across distribution, transportation, and supplier networks. Proven track record of delivering multi-million-dollar cost savings, leading cross-functional transformations, and building data-driven operating models across retail, consulting, and healthcare environments. Advanced capability in KPI governance, supplier performance management, and digital system integration (Excel, SQL, Tableau, SAP), partnering with executive stakeholders to enhance visibility, scalability, and operational efficiency in complex global networks."
},
"bullet_replacements": [
{
"match_anchor": "Directed a $50M supply chain transformation across 15 distribution centers, established program governance,",
"replacement_text": "Directed a $50M end-to-end supply chain transformation across 15 distribution centers, establishing program governance, network standardization, and Lean warehouse optimization, delivering $7.5M in logistics cost reduction within 7 months"
},
{
"match_anchor": "Managed an 8,000+ container network across 15 international warehouses",
"replacement_text": "Managed an 8,000+ container global transportation network within a $240M operation, leading a 35-person cross-functional team and deploying network optimization and freight cost control strategies to reduce annual spend by $400K"
},
{
"match_anchor": "Developed supplier performance dashboards in Tableau by integrating SAP extracts, automating",
"replacement_text": "Built supplier performance management dashboards in Tableau by integrating SAP data extracts, automating KPI reporting, and enabling data-backed contract renegotiations that generated $500K in annual savings"
},
{
"match_anchor": "Conducted 25+ supplier quality audits in 12 months using",
"replacement_text": "Executed 25+ supplier compliance and operational audits within 12 months, implementing corrective action frameworks that increased compliance by 20% and reduced cycle time by 10 days"
},
{
"match_anchor": "Built a go-to-market supply and pricing model by evaluating",
"replacement_text": "Developed a supplier sourcing and margin optimization model by analyzing international cost structures and pricing scenarios, enabling product launch at 60% gross margin"
}
]
}
"""

payload = json.loads(json_payload)

print("=" * 100)
print("ANCHOR MATCHING ANALYSIS")
print("=" * 100)

# Get all paragraph texts from original
original_texts = [para.text for para in original.paragraphs if para.text.strip()]

print("\n1. SUMMARY REPLACEMENT ANCHOR:")
print("-" * 100)
summary_anchor = payload["summary_replacement"]["match_anchor"]
print(f"Anchor: '{summary_anchor}'")

found = False
for idx, text in enumerate(original_texts):
    if summary_anchor in text:
        print(f"✅ FOUND in paragraph {idx}")
        print(f"Full text: '{text}'")
        found = True
        break

if not found:
    print("❌ NOT FOUND - checking for partial matches...")
    for idx, text in enumerate(original_texts):
        if "MBA candidate" in text:
            print(f"Found similar paragraph {idx}: '{text[:100]}...'")

print("\n2. BULLET REPLACEMENT ANCHORS:")
print("-" * 100)

for i, bullet in enumerate(payload["bullet_replacements"]):
    anchor = bullet["match_anchor"]
    print(f"\nBullet {i+1} Anchor: '{anchor}'")
    
    found = False
    for idx, text in enumerate(original_texts):
        if anchor in text:
            print(f"✅ FOUND in paragraph {idx}")
            print(f"Full text: '{text}'")
            found = True
            break
    
    if not found:
        print("❌ NOT FOUND - checking for close matches...")
        # Check for first few words
        first_words = ' '.join(anchor.split()[:5])
        for idx, text in enumerate(original_texts):
            if first_words in text:
                print(f"Similar paragraph {idx}: '{text[:120]}...'")

print("\n" + "=" * 100)
print("OPTIMIZED DOCUMENT - WHAT'S THERE NOW?")
print("=" * 100)

optimized_texts = [para.text for para in optimized.paragraphs if para.text.strip()]

print("\nSummary paragraph:")
for idx, text in enumerate(optimized_texts):
    if "MBA" in text and len(text) > 100:
        print(f"Para {idx}: '{text[:150]}...'")
        break

print("\nBullet paragraphs (containing 'Directed', 'Managed', 'Developed', 'Conducted', 'Built'):")
for idx, text in enumerate(optimized_texts):
    if any(word in text for word in ["Directed a $50M", "Managed an 8,000", "Developed supplier", "Built supplier", "Conducted 25+", "Executed 25+", "Built a go-to-market"]):
        print(f"\nPara {idx}: '{text}'")
