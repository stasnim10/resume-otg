import io
import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import mcp_resume_service as service


@pytest.fixture
def saved_resume(monkeypatch, tmp_path):
    doc = Document()
    doc.add_paragraph("Experienced Python developer.").runs[0].bold = True
    doc.add_paragraph("Built reliable reporting tools.")
    stream = io.BytesIO()
    doc.save(stream)
    original = stream.getvalue()
    monkeypatch.setattr(service, "read_default_resume", lambda: ({"original_name": "resume.docx"}, original))
    monkeypatch.setattr(service, "OUTPUT_DIR", tmp_path)
    return original


def test_prepare_export_download_preserves_source_and_format(saved_resume):
    context = service.prepare_resume("Seeking a Python developer to build reliable reporting systems.")
    payload = {"summary_replacement": {"match_anchor": context["paragraphs"][0], "replacement_text": "Python developer experienced in reporting tools."}}
    result = service.export_resume(payload, context["resume_sha256"])
    doc = Document(io.BytesIO(service.read_export(result["export_id"])))
    assert doc.paragraphs[0].text == payload["summary_replacement"]["replacement_text"]
    assert doc.paragraphs[0].runs[0].bold
    assert doc.paragraphs[1].text == "Built reliable reporting tools."
    assert service.read_default_resume()[1] == saved_resume


def test_invalid_anchor_is_atomic(saved_resume):
    context = service.prepare_resume("Seeking a Python developer to build reliable reporting systems.")
    payload = {"bullet_replacements": [
        {"match_anchor": context["paragraphs"][0], "replacement_text": "Valid replacement"},
        {"match_anchor": "Missing paragraph", "replacement_text": "Invalid replacement"},
    ]}
    with pytest.raises(ValueError, match="uniquely"):
        service.export_resume(payload, context["resume_sha256"])
    assert list(service.OUTPUT_DIR.iterdir()) == []


def test_stale_resume_and_path_traversal_rejected(saved_resume):
    with pytest.raises(ValueError, match="changed"):
        service.export_resume({"bullet_replacements": []}, "stale")
    with pytest.raises(ValueError, match="Invalid export"):
        service.read_export("../../private")


def test_private_job_url_rejected(monkeypatch):
    monkeypatch.setattr("socket.getaddrinfo", lambda *args: [(2, 1, 6, "", ("127.0.0.1", 80))])
    with pytest.raises(ValueError, match="public job link"):
        service.validate_job_url("http://localhost/job")


def test_url_input_uses_safe_fetcher(saved_resume, monkeypatch):
    def fetch(url, **kwargs):
        assert kwargs["validate_url"] is service.validate_job_url
        return "Seeking a Python developer to build reliable reporting systems.", url, "Developer"
    monkeypatch.setattr(service, "fetch_job_description_from_url", fetch)
    assert service.prepare_resume("https://example.com/job")["source_url"] == "https://example.com/job"
