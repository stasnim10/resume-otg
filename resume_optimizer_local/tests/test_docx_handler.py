from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from docx_handler import (
    apply_replacements,
    extract_text,
    replace_exact_paragraph,
    replace_paragraph_text,
)


def _save_doc(path: Path, paragraphs: list[str]) -> str:
    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    doc.save(path)
    return str(path)


def _load_paragraphs(path: Path) -> list[str]:
    doc = Document(path)
    return [p.text for p in doc.paragraphs]


def test_extract_text_preserves_paragraph_order(tmp_path: Path) -> None:
    path = tmp_path / "ordered.docx"
    _save_doc(path, ["First", "Second", "Third"])
    assert extract_text(str(path)).splitlines() == ["First", "Second", "Third"]


def test_extract_text_joins_with_newlines(tmp_path: Path) -> None:
    path = tmp_path / "newlines.docx"
    _save_doc(path, ["A", "B"])
    assert extract_text(str(path)) == "A\nB"


def test_extract_text_keeps_empty_paragraph_slots(tmp_path: Path) -> None:
    path = tmp_path / "empty-lines.docx"
    _save_doc(path, ["Alpha", "", "Gamma"])
    assert extract_text(str(path)).splitlines() == ["Alpha", "", "Gamma"]


def test_replace_paragraph_text_replaces_content() -> None:
    doc = Document()
    paragraph = doc.add_paragraph("Old value")
    assert replace_paragraph_text(paragraph, "New value") is True
    assert paragraph.text == "New value"


def test_replace_paragraph_text_preserves_bold_from_first_run() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("Old ").bold = True
    paragraph.add_run("value").bold = False
    replace_paragraph_text(paragraph, "Replaced")
    assert paragraph.runs[0].bold is True
    assert paragraph.text == "Replaced"


def test_replace_paragraph_text_handles_no_runs_edge_case() -> None:
    doc = Document()
    paragraph = doc.add_paragraph()
    assert not paragraph.runs
    replace_paragraph_text(paragraph, "Inserted")
    assert paragraph.text == "Inserted"


def test_replace_exact_paragraph_successful_match() -> None:
    doc = Document()
    doc.add_paragraph("Anchor text")
    message = replace_exact_paragraph(doc, "Anchor text", "New text")
    assert "✅ Replaced" in message
    assert doc.paragraphs[0].text == "New text"


def test_replace_exact_paragraph_raises_when_no_match() -> None:
    doc = Document()
    doc.add_paragraph("Different text")
    with pytest.raises(ValueError, match="Anchor not found"):
        replace_exact_paragraph(doc, "Anchor text", "New text")


def test_replace_exact_paragraph_raises_on_duplicate_match() -> None:
    doc = Document()
    doc.add_paragraph("Anchor text")
    doc.add_paragraph("Anchor text")
    with pytest.raises(ValueError, match="Multiple matches found"):
        replace_exact_paragraph(doc, "Anchor text", "New text")


def test_replace_exact_paragraph_trims_whitespace_for_matching() -> None:
    doc = Document()
    doc.add_paragraph("Anchor text")
    replace_exact_paragraph(doc, "  Anchor text  ", "New text")
    assert doc.paragraphs[0].text == "New text"


def test_apply_replacements_success_for_summary_only(tmp_path: Path) -> None:
    path = tmp_path / "summary.docx"
    _save_doc(path, ["Summary old", "Other"])
    payload = {"summary_replacement": {"match_anchor": "Summary old", "replacement_text": "Summary new"}}
    success, message = apply_replacements(str(path), payload)
    assert success is True
    assert "Replaced 1 section(s)." in message
    assert _load_paragraphs(tmp_path / "summary_Optimized.docx")[0] == "Summary new"


def test_apply_replacements_success_for_bullet_only(tmp_path: Path) -> None:
    path = tmp_path / "bullet.docx"
    _save_doc(path, ["Bullet old"])
    payload = {"bullet_replacements": [{"match_anchor": "Bullet old", "replacement_text": "Bullet new"}]}
    success, _ = apply_replacements(str(path), payload)
    assert success is True
    assert _load_paragraphs(tmp_path / "bullet_Optimized.docx")[0] == "Bullet new"


def test_apply_replacements_success_for_skills_only(tmp_path: Path) -> None:
    path = tmp_path / "skills.docx"
    _save_doc(path, ["Skills old"])
    payload = {"skills_replacements": [{"match_anchor": "Skills old", "replacement_text": "Skills new"}]}
    success, _ = apply_replacements(str(path), payload)
    assert success is True
    assert _load_paragraphs(tmp_path / "skills_Optimized.docx")[0] == "Skills new"


def test_apply_replacements_success_for_multi_section_payload(tmp_path: Path) -> None:
    path = tmp_path / "multi.docx"
    _save_doc(path, ["Summary old", "Bullet old", "Skills old"])
    payload = {
        "summary_replacement": {"match_anchor": "Summary old", "replacement_text": "Summary new"},
        "bullet_replacements": [{"match_anchor": "Bullet old", "replacement_text": "Bullet new"}],
        "skills_replacements": [{"match_anchor": "Skills old", "replacement_text": "Skills new"}],
    }
    success, message = apply_replacements(str(path), payload)
    assert success is True
    assert "Replaced 3 section(s)." in message
    assert _load_paragraphs(tmp_path / "multi_Optimized.docx") == ["Summary new", "Bullet new", "Skills new"]


def test_apply_replacements_fails_when_anchor_missing(tmp_path: Path) -> None:
    path = tmp_path / "missing.docx"
    _save_doc(path, ["Present text"])
    payload = {"summary_replacement": {"match_anchor": "Missing text", "replacement_text": "New"}}
    success, message = apply_replacements(str(path), payload)
    assert success is False
    assert "Anchor not found" in message


def test_apply_replacements_fails_when_duplicate_anchor(tmp_path: Path) -> None:
    path = tmp_path / "duplicate.docx"
    _save_doc(path, ["Dup", "Dup"])
    payload = {"summary_replacement": {"match_anchor": "Dup", "replacement_text": "New"}}
    success, message = apply_replacements(str(path), payload)
    assert success is False
    assert "Multiple matches found" in message


def test_apply_replacements_continues_collecting_errors(tmp_path: Path) -> None:
    path = tmp_path / "errors.docx"
    _save_doc(path, ["One", "Two"])
    payload = {
        "bullet_replacements": [
            {"match_anchor": "Missing one", "replacement_text": "A"},
            {"match_anchor": "Missing two", "replacement_text": "B"},
        ]
    }
    success, message = apply_replacements(str(path), payload)
    assert success is False
    assert message.count("Anchor not found") == 2


def test_apply_replacements_does_not_create_output_file_on_failure(tmp_path: Path) -> None:
    path = tmp_path / "no-output-on-fail.docx"
    _save_doc(path, ["Source"])
    payload = {"summary_replacement": {"match_anchor": "Missing", "replacement_text": "New"}}
    success, _ = apply_replacements(str(path), payload)
    assert success is False
    assert not (tmp_path / "no-output-on-fail_Optimized.docx").exists()


def test_apply_replacements_returns_global_error_for_invalid_doc_path(tmp_path: Path) -> None:
    payload = {"summary_replacement": {"match_anchor": "A", "replacement_text": "B"}}
    success, message = apply_replacements(str(tmp_path / "missing.docx"), payload)
    assert success is False
    assert "❌ Error:" in message
