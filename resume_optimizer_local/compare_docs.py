from docx import Document

# Load both documents
original = Document('/Users/simum/Personal/My Mac/Documents/Simon Business School/Job Hunt/2026 Full Time Roles/USA/Draft/Simum Tasnim_Resume_Draft.docx')
optimized = Document('/Users/simum/Personal/My Mac/Documents/Simon Business School/Job Hunt/2026 Full Time Roles/USA/Draft/Simum Tasnim_Resume_Draft_Optimized.docx')

print("=" * 90)
print("ORIGINAL DOCUMENT")
print("=" * 90)
print(f"Total paragraphs: {len(original.paragraphs)}\n")

for i, para in enumerate(original.paragraphs[:50]):
    if para.text.strip():
        style = para.style.name if para.style else "No Style"
        text = para.text[:70]
        print(f"{i:2d}. [{style:20s}] {text}")

print("\n" + "=" * 90)
print("OPTIMIZED DOCUMENT")
print("=" * 90)
print(f"Total paragraphs: {len(optimized.paragraphs)}\n")

for i, para in enumerate(optimized.paragraphs[:50]):
    if para.text.strip():
        style = para.style.name if para.style else "No Style"
        text = para.text[:70]
        print(f"{i:2d}. [{style:20s}] {text}")
